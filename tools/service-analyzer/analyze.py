"""Analyse an order-of-service PDF against its hand-made PPTX.

Usage: analyze.py <order.pdf> <deck.pptx> <out_dir>
Writes <out_dir>/analysis.json and <out_dir>/report.md.
"""
import collections
import difflib
import json
import re
import statistics
import sys

import fitz
from pptx import Presentation
from pptx.util import Emu


def words(text):
    text = text.lower().replace("’", "'").replace("‘", "'")
    return re.findall(r"[\w']+", text)


def span_style(span):
    font = span["font"].lower()
    bold = bool(span["flags"] & 16) or "bold" in font or "black" in font
    italic = bool(span["flags"] & 2) or "italic" in font or "oblique" in font
    return bold, italic


def extract_pdf(path):
    doc = fitz.open(path)
    lines = []
    page_info = []
    for pno, page in enumerate(doc, 1):
        page_info.append({"page": pno, "images": len(page.get_images()), "w": page.rect.width, "h": page.rect.height})
        for block in page.get_text("dict")["blocks"]:
            if block["type"] != 0:
                continue
            for line in block["lines"]:
                spans = [s for s in line["spans"] if s["text"].strip()]
                if not spans:
                    continue
                text = "".join(s["text"] for s in line["spans"]).strip()
                styles = [span_style(s) for s in spans]
                lines.append({
                    "page": pno,
                    "text": text,
                    "size": round(max(s["size"] for s in spans), 1),
                    "bold": all(b for b, _ in styles),
                    "italic": all(i for _, i in styles),
                    "font": spans[0]["font"],
                    "color": spans[0]["color"],
                    "x0": round(line["bbox"][0], 1),
                    "y0": round(line["bbox"][1], 1),
                })
    char_by_size = collections.Counter()
    for ln in lines:
        char_by_size[ln["size"]] += len(ln["text"])
    body = char_by_size.most_common(1)[0][0]
    for ln in lines:
        if ln["bold"] and ln["size"] >= body + 1.5:
            ln["role"] = "heading"
        elif ln["size"] <= body - 1.5:
            ln["role"] = "small"
        elif ln["italic"]:
            ln["role"] = "rubric"
        elif ln["bold"]:
            ln["role"] = "response"
        else:
            ln["role"] = "body"
    return lines, page_info, body, char_by_size


def eff_size(run, para, shape):
    for obj in (run.font, para.font):
        if obj.size:
            return obj.size.pt
    return None


def extract_pptx(path):
    prs = Presentation(path)
    sw, sh = prs.slide_width, prs.slide_height
    slides = []
    for idx, slide in enumerate(prs.slides, 1):
        shapes = []
        all_text = []
        sizes = []
        lines = 0
        max_chars = 0
        for shape in sorted(slide.shapes, key=lambda s: (round((s.top or 0) / 300000), s.left or 0)):
            entry = {
                "name": shape.name,
                "kind": str(shape.shape_type),
                "box": [round(v / d, 3) for v, d in
                        ((shape.left or 0, sw), (shape.top or 0, sh), (shape.width or 0, sw), (shape.height or 0, sh))],
            }
            if shape.has_text_frame:
                tf = shape.text_frame
                entry["autofit"] = str(tf.auto_size)
                paras = []
                for p in tf.paragraphs:
                    ptext = p.text.replace("\x0b", "\n")
                    if not ptext.strip():
                        continue
                    run_info = []
                    for r in p.runs:
                        if not r.text.strip():
                            continue
                        sz = eff_size(r, p, shape)
                        if sz:
                            sizes.append(sz)
                        run_info.append({"size": sz, "bold": r.font.bold, "italic": r.font.italic})
                    for seg in re.split(r"[\n\x0b]", ptext):
                        if seg.strip():
                            lines += 1
                            max_chars = max(max_chars, len(seg.strip()))
                    paras.append({"text": ptext, "runs": run_info})
                    all_text.append(ptext)
                entry["paragraphs"] = paras
            elif getattr(shape, "has_table", False) and shape.has_table:
                cells = [[c.text.replace("\x0b", "\n") for c in row.cells] for row in shape.table.rows]
                entry["table"] = {"rows": len(cells), "cols": len(cells[0]) if cells else 0}
                for row in cells:
                    for t in row:
                        if t.strip():
                            all_text.append(t)
                            for seg in t.split("\n"):
                                if seg.strip():
                                    lines += 1
                                    max_chars = max(max_chars, len(seg.strip()))
            elif shape.shape_type == 13:
                entry["picture"] = True
            shapes.append(entry)
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        text = "\n".join(all_text)
        slides.append({
            "slide": idx,
            "layout": slide.slide_layout.name,
            "shapes": shapes,
            "text": text,
            "words": len(words(text)),
            "lines": lines,
            "max_line_chars": max_chars,
            "sizes": sorted(set(sizes)),
            "notes": notes,
            "blank": not text.strip() and not any(s.get("picture") for s in shapes),
        })
    return slides, (sw, sh)


def align(pdf_lines, slides):
    a, a_page, a_role = [], [], []
    for ln in pdf_lines:
        for w in words(ln["text"]):
            a.append(w)
            a_page.append(ln["page"])
            a_role.append(ln["role"])
    b, b_slide = [], []
    for s in slides:
        for w in words(s["text"]):
            b.append(w)
            b_slide.append(s["slide"])
    sm = difflib.SequenceMatcher(None, a, b, autojunk=False)
    kept = collections.Counter()
    slide_pages = collections.defaultdict(collections.Counter)
    dropped = collections.defaultdict(list)
    added = collections.defaultdict(list)
    changed = []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            for k in range(i1, i2):
                kept[a_page[k]] += 1
            for k in range(j1, j2):
                slide_pages[b_slide[j1]][a_page[i1]] += 0
            for off in range(i2 - i1):
                slide_pages[b_slide[j1 + off]][a_page[i1 + off]] += 1
        elif tag == "delete":
            if i2 - i1 >= 4:
                dropped[a_page[i1]].append(" ".join(a[i1:i2]))
        elif tag == "insert":
            if j2 - j1 >= 4:
                added[b_slide[j1]].append(" ".join(b[j1:j2]))
        else:
            if (i2 - i1) <= 12 and (j2 - j1) <= 12:
                changed.append({"page": a_page[i1], "slide": b_slide[j1],
                                "pdf": " ".join(a[i1:i2]), "pptx": " ".join(b[j1:j2])})
            else:
                if i2 - i1 >= 4:
                    dropped[a_page[i1]].append(" ".join(a[i1:i2]))
                if j2 - j1 >= 4:
                    added[b_slide[j1]].append(" ".join(b[j1:j2]))
    page_words = collections.Counter(a_page)
    return {
        "pdf_words": len(a),
        "pptx_words": len(b),
        "ratio": round(sm.ratio(), 3),
        "kept_by_page": {p: [kept[p], page_words[p]] for p in sorted(page_words)},
        "slide_to_pages": {s: dict(c.most_common()) for s, c in sorted(slide_pages.items()) if any(c.values())},
        "dropped": {p: v for p, v in sorted(dropped.items())},
        "added": {s: v for s, v in sorted(added.items())},
        "changed": changed,
    }


def section_map(pdf_lines):
    heads = [(ln["page"], ln["text"]) for ln in pdf_lines if ln["role"] == "heading"]
    return heads


def report(pdf_lines, page_info, body, char_by_size, slides, slide_dim, al):
    out = []
    w = out.append
    w("# Service order analysis\n")
    w(f"PDF: {len(page_info)} pages, body size {body} pt. PPTX: {len(slides)} slides, "
      f"{slide_dim[0] / 914400:.2f}x{slide_dim[1] / 914400:.2f} in.\n")
    w("## PDF typography (characters by size)\n")
    for size, n in sorted(char_by_size.items(), reverse=True):
        w(f"- {size} pt: {n} chars")
    roles = collections.Counter(ln["role"] for ln in pdf_lines)
    w(f"\nLine roles: {dict(roles)}\n")
    fonts = collections.Counter(ln["font"] for ln in pdf_lines)
    w(f"Fonts: {dict(fonts)}\n")
    w("## PDF section headings (page: text)\n")
    for p, t in section_map(pdf_lines):
        w(f"- p{p}: {t}")
    w("\n## Slide metrics\n")
    wc = [s["words"] for s in slides if not s["blank"]]
    w(f"Words per slide: median {statistics.median(wc)}, mean {statistics.mean(wc):.1f}, max {max(wc)}, min {min(wc)}")
    lc = [s["lines"] for s in slides if not s["blank"]]
    w(f"Lines per slide: median {statistics.median(lc)}, max {max(lc)}")
    mc = [s["max_line_chars"] for s in slides if not s["blank"]]
    w(f"Longest line chars: median {statistics.median(mc)}, max {max(mc)}")
    size_hist = collections.Counter(sz for s in slides for sz in s["sizes"])
    w(f"Font sizes used (slides using each): {dict(sorted(size_hist.items()))}")
    layouts = collections.Counter(s["layout"] for s in slides)
    w(f"Layouts: {dict(layouts)}")
    autofit = collections.Counter(sh.get("autofit") for s in slides for sh in s["shapes"] if "autofit" in sh)
    w(f"Text autofit modes: {dict(autofit)}")
    w(f"Blank slides: {[s['slide'] for s in slides if s['blank']]}")
    w(f"Slides with notes text: {[s['slide'] for s in slides if s['notes']]}")
    w("\n## Per slide\n")
    w("| slide | words | lines | max chars | sizes | layout | PDF pages (words matched) |")
    w("|---|---|---|---|---|---|---|")
    for s in slides:
        pages = al["slide_to_pages"].get(s["slide"], {})
        pg = ", ".join(f"p{p}:{n}" for p, n in list(pages.items())[:3])
        w(f"| {s['slide']} | {s['words']} | {s['lines']} | {s['max_line_chars']} | {s['sizes']} | {s['layout']} | {pg} |")
    w(f"\n## Alignment (similarity {al['ratio']}; PDF {al['pdf_words']} words, PPTX {al['pptx_words']} words)\n")
    w("### PDF words kept in slides, per page (kept/total)\n")
    for p, (k, t) in al["kept_by_page"].items():
        w(f"- p{p}: {k}/{t} ({100 * k // max(t, 1)}%)")
    w("\n### Wording changes (PDF -> PPTX), short edits only\n")
    for c in al["changed"]:
        w(f"- p{c['page']}/s{c['slide']}: \"{c['pdf']}\" -> \"{c['pptx']}\"")
    w("\n### PDF text not in slides (4+ words), by page\n")
    for p, items in al["dropped"].items():
        for t in items:
            w(f"- p{p}: {t[:160]}")
    w("\n### Slide text not in PDF (4+ words), by slide\n")
    for s, items in al["added"].items():
        for t in items:
            w(f"- s{s}: {t[:160]}")
    return "\n".join(out)


def main():
    pdf_path, pptx_path, out_dir = sys.argv[1:4]
    pdf_lines, page_info, body, char_by_size = extract_pdf(pdf_path)
    slides, dim = extract_pptx(pptx_path)
    al = align(pdf_lines, slides)
    with open(f"{out_dir}/analysis.json", "w", encoding="utf-8") as f:
        json.dump({"pdf_lines": pdf_lines, "pages": page_info, "slides": slides, "alignment": al}, f,
                  ensure_ascii=False, indent=1, default=str)
    with open(f"{out_dir}/report.md", "w", encoding="utf-8") as f:
        f.write(report(pdf_lines, page_info, body, char_by_size, slides, dim, al))
    print("ok", len(pdf_lines), "pdf lines,", len(slides), "slides")


if __name__ == "__main__":
    main()
